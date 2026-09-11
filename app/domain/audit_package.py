"""The auditor's export.

`package.py` builds the rate proposal workbook, which needs a sealed rate. This
builds the thing that exists before one: the record. Controls and whether they
tie, every judgment with its reasoning and citation, what was split and why,
what evidence is attached and to what, who signed for their own effort and in
what words, everything that happened, and what is still open.

The point is that an auditor can take it away. A system that only answers
questions on screen makes the reviewer work at your desk, on your schedule.
"""

from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HEAD_FILL = PatternFill("solid", fgColor="E8EDEC")
WRAP = Alignment(wrap_text=True, vertical="top")
MONEY = "#,##0.00"


def _sheet(wb: Workbook, name: str, title: str, subtitle: str = ""):
    ws = wb.create_sheet(name[:31])
    ws["A1"] = title
    ws["A1"].font = TITLE
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True, size=9)
    ws.freeze_panes = "A5"
    return ws


def _table(ws, row: int, headers: list[str], rows: list[dict],
           keys: list[str], widths: list[int] | None = None,
           money_cols: set[int] = frozenset()) -> int:
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = WRAP
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = row
    for item in rows:
        r += 1
        for i, key in enumerate(keys, start=1):
            value = item.get(key)
            if isinstance(value, Decimal):
                value = float(value)
            elif isinstance(value, datetime):
                value = value.replace(tzinfo=None)
            elif isinstance(value, (dict, list)):
                value = str(value)
            cell = ws.cell(row=r, column=i, value=value)
            cell.alignment = WRAP
            if i in money_cols:
                cell.number_format = MONEY
    return r


def build_audit_package(*, period: str, out_path: Path, controls: list[dict],
                        decisions: list[dict], segments: list[dict],
                        evidence: list[dict], certifications: list[dict],
                        activity: list[dict], worklist: list[dict],
                        rollup: dict, generated_by: str,
                        exceptions: list[dict] | None = None,
                        materiality: list[dict] | None = None,
                        rates: list[dict] | None = None,
                        allocations: list[dict] | None = None,
                        gl_pl: list[dict] | None = None,
                        gl_bs: list[dict] | None = None,
                        reconciling_items: list[dict] | None = None) -> Path:
    wb = Workbook()

    # ── Index ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Index"
    ws["A1"] = f"Youngstown Business Incubator — {period} cost record"
    ws["A1"].font = TITLE
    ws["A2"] = (f"Generated {datetime.now(timezone.utc):%d %b %Y %H:%M} UTC "
                f"by {generated_by}")
    ws["A2"].font = Font(italic=True, size=9)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 78

    rows = [
        ("Controls", "Every derived figure and whether it ties to its source"),
        ("A-1 Reconciliation",
         "The general ledger against the profit and loss and the balance "
         "sheet, and every difference between them by name"),
        ("Decisions", "Each classification, its reasoning and its citation"),
        ("Segments", "Mixed lines split into parts, and what each part is"),
        ("Evidence", "Documents on file and what they support"),
        ("Certifications", "Effort attested, by whom, in what words"),
        ("Activity", "Everything that happened, in order"),
        ("Open items", "What is not finished, and why it matters"),
    ]
    r = 4
    ws.cell(row=r, column=1, value="Sheet").font = BOLD
    ws.cell(row=r, column=2, value="What it shows").font = BOLD
    for name, desc in rows:
        r += 1
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=desc).alignment = WRAP

    r += 2
    ws.cell(row=r, column=1, value="Where it stands").font = BOLD
    for label, key, fmt in (("Income", "income", MONEY),
                            ("Expenses", "expense", MONEY),
                            ("Cost of goods", "cogs", MONEY),
                            ("Other income", "other_income", MONEY),
                            ("Net income", "net_income", MONEY),
                            ("Ledger lines", "ledger_lines", None),
                            ("Lines in P&L scope", "scope_lines", None),
                            ("Decisions recorded", "decisions", None),
                            ("Documents on file", "documents", None)):
        r += 1
        ws.cell(row=r, column=1, value=label)
        v = rollup.get(key)
        cell = ws.cell(row=r, column=2,
                       value=float(v) if isinstance(v, Decimal) else v)
        if fmt:
            cell.number_format = fmt

    r += 2
    ws.cell(row=r, column=1, value="Note").font = BOLD
    note = ws.cell(row=r, column=2, value=(
        "Unclassified cost is never defaulted into a pool. While the open "
        "items sheet is long the indirect rate reads high, which is the "
        "honest direction to err. No rate may be computed until the decision "
        "set is sealed, and no rate survives a classification changing "
        "underneath it."))
    note.alignment = WRAP

    # ── Controls ─────────────────────────────────────────────────────
    ws = _sheet(wb, "Controls", "Controls",
                "A derived figure that cannot be reconciled to its source does "
                "not ship.")
    # State, not just Ties. A control that could not be evaluated reads as a
    # variance without it: with no asset register on file, ASSET_REGISTER
    # showed -850,382.89 and Ties=False, which a reviewer reads as "the
    # register disagrees with the ledger by the whole of its depreciation"
    # rather than "there is no register". Those are different findings and
    # only one of them is true. The register already carries the distinction;
    # the workpaper dropped it, which is the only place it matters.
    _table(ws, 5,
           ["Control", "What it proves", "Variance or exceptions", "State",
            "Ties"],
           controls, ["control", "description", "variance", "state", "ties"],
           [22, 54, 20, 12, 10], {3})

    # ── A-1 Reconciliation ───────────────────────────────────────────
    #
    # Ahead of the decisions, because it is what the decisions rest on. A
    # reviewer who does not believe the ledger is the books has no reason to
    # read anything after it.
    ws = _sheet(wb, "A-1 Reconciliation", "Cross-reference reconciliation",
                "The three source documents against each other. Run before "
                "anything was classified, which is the only order in which a "
                "reconciliation is worth anything.")
    r = _table(ws, 5,
               ["Control", "What it proves", "", "Ledger side", "",
                "Statement side", "Variance or exceptions", "State", "Ties"],
               controls,
               ["control", "description", "left_label", "left_value",
                "right_label", "right_value", "variance", "state", "ties"],
               # The sheet stacks four tables; these widths have to
               # suit all of them, because _table rewrites the column
               # dimensions and the later calls leave them alone.
               [50, 46, 26, 16, 26, 16, 20, 12, 14, 12], {4, 6, 7})

    if gl_pl:
        r += 2
        ws.cell(row=r, column=1, value="Where the ledger and the P&L differ").font = BOLD
        r += 1
        ws.cell(row=r, column=1, value=(
            "Sections tie while these accounts do not, because money moved "
            "between two expense accounts nets to nothing at the section "
            "line. Each row is named below or it is unexplained.")).alignment = WRAP
        r = _table(ws, r + 2,
                   ["Account", "Section", "Ledger", "Lines", "P&L",
                    "Named", "Difference", "Unexplained"],
                   gl_pl,
                   ["account", "section", "gl_amount", "gl_lines", "pl_amount",
                    "reconciling", "gross_variance", "unexplained"],
                   None, {3, 5, 6, 7, 8})

    if reconciling_items:
        r += 2
        ws.cell(row=r, column=1, value="Reconciling items").font = BOLD
        r += 1
        ws.cell(row=r, column=1, value=(
            "Each one names the ledger lines it consists of, and the database "
            "refuses it unless those lines add to the amount claimed. That is "
            "what separates a reconciling item from a plug.")).alignment = WRAP
        r = _table(ws, r + 2,
                   ["Ledger puts it in", "The statement puts it in", "Amount",
                    "Lines", "Kind", "Recorded by", "Explanation"],
                   reconciling_items,
                   ["from_account", "to_account", "amount", "lines", "kind",
                    "recorded_by", "explanation"],
                   None, {3})

    if gl_bs:
        r += 2
        ws.cell(row=r, column=1,
                value="Balance sheet accounts the ledger does not simply confirm").font = BOLD
        r += 1
        ws.cell(row=r, column=1, value=(
            "Opening balance plus the year's movement against the sheet as "
            "printed. An account the sheet omits has to close at zero here, "
            "which is checked rather than assumed.")).alignment = WRAP
        _table(ws, r + 2,
               ["Ledger account", "As the sheet names it", "Opening",
                "Movement", "Closing", "Balance sheet", "Variance",
                "Closed at zero", "By alias"],
               gl_bs,
               ["account", "bs_leaf", "opening", "activity", "closing",
                "bs_amount", "variance", "absent_because_zero",
                "matched_by_alias"],
               None, {3, 4, 5, 6, 7})

    # ── Decisions ────────────────────────────────────────────────────
    ws = _sheet(wb, "Decisions", "Classification decisions",
                "Each judgment, who made it, the reasoning, and the authority "
                "relied on.")
    _table(ws, 5,
           ["Scope", "Pool", "990 function", "Federal", "Grade", "Amount",
            "Rationale", "Citation", "Decided by", "Decided at", "Reversed"],
           decisions,
           ["scope", "pool", "function_990", "federal", "grade", "amount",
            "rationale", "citation", "decided_by", "decided_at", "reversed_at"],
           [40, 13, 20, 13, 24, 14, 62, 20, 18, 18, 18], {6})

    # ── Segments ─────────────────────────────────────────────────────
    ws = _sheet(wb, "Segments", "Segmented lines",
                "A booked entry is often not one thing. Each split reconciles "
                "to its source line to the cent.")
    _table(ws, 5,
           ["Batch", "Label", "Segments", "Amount", "Rationale", "Citation",
            "Created by", "Created at", "Reversed"],
           segments,
           ["batch_key", "label", "segments", "amount", "rationale",
            "citation", "created_by", "created_at", "reversed_at"],
           [22, 40, 11, 14, 62, 20, 18, 18, 18], {4})

    # ── Evidence ─────────────────────────────────────────────────────
    ws = _sheet(wb, "Evidence", "Evidence on file",
                "Content addressed by SHA-256, so the document produced later "
                "is provably the one relied on.")
    _table(ws, 5,
           ["Evidence", "Kind", "SHA-256", "Bytes", "Attached to", "Relevance",
            "Received from", "Received at"],
           evidence,
           ["evidence_id", "kind", "sha256", "byte_size", "attachments",
            "relevance", "received_from", "received_at"],
           [20, 22, 68, 12, 13, 56, 20, 20])

    # ── Certifications ───────────────────────────────────────────────
    ws = _sheet(wb, "Certifications", "Effort certifications",
                "2 CFR 200.430(i). The words signed are reproduced in full; a "
                "certification whose wording cannot be produced is not one.")
    _table(ws, 5,
           ["Employee", "Role", "Signed by", "Signed at", "Objectives",
            "Statement", "Distribution as signed", "Superseded"],
           certifications,
           ["employee_key", "certifier_role", "signed_by", "signed_at",
            "objectives", "statement", "distribution", "superseded_at"],
           [16, 13, 20, 20, 12, 70, 70, 20])

    # ── Rates ────────────────────────────────────────────────────────
    if rates:
        ws = _sheet(wb, "Rates", "Rates and allocation",
                    "Each rate carries the seal of the decision set it came "
                    "from. A rate whose seal does not match a sealed set is "
                    "refused by the database, not by a handler.")
        r = _table(ws, 5, ["Kind", "Pool", "Base type", "Base", "Rate",
                           "Status", "Seal", "Computed by", "When"],
                   rates,
                   ["kind", "pool_amount", "base_type", "base_amount", "rate",
                    "status", "seal_hash", "computed_by", "computed_at"],
                   [20, 16, 18, 16, 12, 14, 20, 22, 20], {2, 4})
        if allocations:
            r += 2
            ws.cell(row=r, column=1, value="Allocation").font = BOLD
            _table(ws, r + 1,
                   ["Rate", "Objective", "Base", "Allocated"], allocations,
                   ["kind", "objective_id", "base_amount", "allocated"],
                   [20, 26, 16, 16], {3, 4})

    # ── Exceptions ───────────────────────────────────────────────────
    #
    # The first schedule a reviewer asks for, and the one that decides whether
    # the rest of the file reads as candid or as managed. A defensible record
    # is not one with no exceptions; it is one where every exception is
    # stated, reasoned, and findable without reading the log.
    if exceptions is not None:
        ws = _sheet(wb, "Exceptions", "Exceptions",
                    "Every place the standard was bent, who bent it and why. "
                    "An exception without a reason is the finding.")
        _table(ws, 5,
               ["When", "Kind", "Who", "Subject", "What", "Why", "Amount"],
               exceptions,
               ["occurred_at", "kind", "actor", "subject", "detail", "reason",
                "amount"],
               [20, 26, 22, 34, 40, 56, 16], {7})

    # ── Materiality ──────────────────────────────────────────────────
    if materiality is not None:
        ws = _sheet(wb, "Materiality", "Evidence against materiality",
                    "The standard each judgment has to meet is a written "
                    "policy applied to its size and its federal exposure, so "
                    "that less work on small items is a position rather than "
                    "an omission.")
        _table(ws, 5,
               ["Scope", "Pool", "Amount", "Grade", "Required", "Meets",
                "Federal", "Decided by"],
               materiality,
               ["scope", "pool", "amount", "grade", "required_grade",
                "meets_standard", "federal", "decided_by"],
               [44, 14, 16, 26, 26, 10, 14, 22], {3})

    # ── Activity ─────────────────────────────────────────────────────
    ws = _sheet(wb, "Activity", "Everything that happened",
                "Read from the append-only records themselves, so the trail "
                "cannot disagree with what it describes.")
    _table(ws, 5,
           ["When", "What", "Who", "Object", "Reference", "Detail", "Amount"],
           activity,
           ["occurred_at", "kind", "actor", "entity", "label", "detail",
            "amount"],
           [20, 18, 20, 16, 46, 46, 14], {7})

    # ── Open items ───────────────────────────────────────────────────
    ws = _sheet(wb, "Open items", "What is not finished",
                "Ordered by money. Blocking items prevent a rate from being "
                "sealed.")
    _table(ws, 5, ["Class", "Severity", "Item", "Detail", "Amount"], worklist,
           ["kind", "severity", "label", "detail", "amount"],
           [26, 12, 54, 52, 14], {5})

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
