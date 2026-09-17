"""The three workbooks that leave the building.

Each one is a deliverable somebody signs: the rate a sponsor is asked to
accept, the report an auditor works from, and the functional allocation the
return prints. They share the audit package's sheet helpers so the whole set
reads as one body of work rather than three exports written by three people.

Every one of them states what is unfinished on its first sheet. A workbook
that looks complete and is not is worse than no workbook, because it travels
— somebody forwards it, somebody quotes a figure out of it, and the caveat
that lived on a screen does not go with it.
"""

from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.domain.audit_package import (BOLD, TITLE, WRAP, certification_lines,
                                      _sheet, _table)

RULE = Font(italic=True, size=9)


def _stamp(ws, period: str, row: int = 3) -> None:
    ws.cell(row=row, column=1,
            value=(f"Period {period} · prepared "
                   f"{datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC")).font = RULE


def _caveat(ws, row: int, lines: list[str]) -> int:
    """What is not finished, at the top, before any figure.

    Deliberately above the numbers rather than in a footnote. A reader who
    has already read the total has already formed a view.
    """
    for text in lines:
        c = ws.cell(row=row, column=1, value=text)
        c.alignment = WRAP
        c.font = (BOLD if text.startswith(("NOT ", "CERTIFIED"))
                  else Font(size=10))
        ws.row_dimensions[row].height = 28
        row += 1
    return row + 1



def _as_rate(ws, first_row: int, last_row: int, col: int) -> None:
    """A rate is read as a percentage to four places, not rounded to cents.

    _table money-formats with #,##0.00, which turned 0.914400 into 0.91 — a
    rate reported to two decimal places of a ratio is a rate nobody can
    reproduce, and 0.91 on a workpaper reads as ninety-one cents.
    """
    for r in range(first_row, last_row + 1):
        c = ws.cell(row=r, column=col)
        if isinstance(c.value, (int, float)):
            c.number_format = "0.0000%"


# ------------------------------------------------------------------ rate


def build_rate_buildup(*, period: str, out_path: Path, rates: list[dict],
                       carve_outs: list[dict], coverage: dict | None,
                       open_controls: list[dict], final: bool,
                       certification: dict | None = None) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    ws = _sheet(wb, "Rate", "Indirect rate build-up",
                "The pool, what was taken out of it and under what authority, "
                "the base, and the seal the whole thing hangs off.")
    _stamp(ws, period)
    pct = Decimal(str((coverage or {}).get("pct_dollars_covered") or 0))
    # The signature first, because it is the question a reader asks of a rate
    # before they ask how complete it is.
    r = _caveat(ws, 5, certification_lines(certification))
    if not final:
        why = []
        if open_controls:
            why.append(f"{len(open_controls)} cross-reference point(s) are open: "
                       + ", ".join(c["control"] for c in open_controls))
        if pct < 100:
            why.append(f"classification is {pct}% complete by dollars — cost "
                       f"nobody has judged is not in any pool, so the rate "
                       f"reads high")
        r = _caveat(ws, r, ["NOT FINAL — this rate is a working figure.", *why])

    head = r
    r = _table(ws, r,
               ["Rate", "Pool", "Base", "Base amount", "Rate", "Status",
                "Sealed by", "Sealed at", "Seal"],
               rates,
               ["kind", "pool_amount", "base_type", "base_amount", "rate",
                "status", "sealed_by", "sealed_at", "seal_hash"],
               [20, 16, 18, 16, 12, 14, 20, 22, 66], {2, 4})
    _as_rate(ws, head + 1, r, 5)

    if carve_outs:
        r += 2
        ws.cell(row=r, column=1, value="Carve-outs").font = BOLD
        r += 1
        ws.cell(row=r, column=1, value=(
            "What was removed from a pool before the rate was struck. Each "
            "one names the authority it was removed under and the driver that "
            "sized it — a reduction without both is an adjustment, not a "
            "carve-out.")).alignment = WRAP
        r = _table(ws, r + 2,
                   ["Pool", "What", "Authority", "Amount", "Driver", "Grade",
                    "Recorded by"],
                   carve_outs,
                   ["pool", "name", "citation", "amount", "driver", "grade",
                    "created_by"],
                   [16, 30, 26, 16, 40, 22, 20], {4})

    wb.save(out_path)
    return out_path


# -------------------------------------------------------- auditor's report


def build_auditors_report(*, period: str, out_path: Path, controls: list[dict],
                          asset_control: dict | None, exceptions: list[dict],
                          coverage: dict | None, evidence_coverage: list[dict],
                          rates: list[dict],
                          reconciling_items: list[dict],
                          certification: dict | None = None) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    # ── Opinion-bearing summary ──────────────────────────────────────
    ws = _sheet(wb, "Report", "Auditor's report",
                "What the engagement asserts, and what proves each assertion.")
    _stamp(ws, period)
    open_controls = [c for c in controls if not c.get("ties")]
    pct = Decimal(str((coverage or {}).get("pct_dollars_covered") or 0))
    r = _caveat(ws, 5, certification_lines(certification))
    r = _caveat(ws, r, [
        ("All cross-reference points tie."
         if not open_controls else
         f"NOT FINAL — {len(open_controls)} cross-reference point(s) open: "
         + ", ".join(c["control"] for c in open_controls)),
        f"Classification is {pct}% complete by dollars. Unclassified cost sits "
        f"in no pool and is never defaulted into one, so any rate below reads "
        f"high while the queue is open — which is the honest direction to err.",
        (f"Asset register: {asset_control.get('state')}"
         + (f" — needs {asset_control.get('needs')}."
            if asset_control.get("needs") else ".")
         ) if asset_control else "",
    ])

    r = _table(ws, r,
               ["Control", "What it proves", "Ledger side", "", "Statement side",
                "", "Variance", "State", "Ties"],
               controls,
               ["control", "description", "left_label", "left_value",
                "right_label", "right_value", "variance", "state", "ties"],
               [24, 50, 24, 16, 24, 16, 16, 12, 10], {4, 6, 7})

    # ── Reconciling items ────────────────────────────────────────────
    ws = _sheet(wb, "Reconciling items", "Differences, named",
                "Each carries the ledger lines it consists of. A deferred "
                "trigger refuses one whose lines do not add to the amount "
                "claimed, which is what separates an item from a plug.")
    _stamp(ws, period)
    _table(ws, 5,
           ["Control", "Ledger puts it in", "The statement puts it in",
            "Amount", "Lines", "Kind", "Recorded by", "Explanation"],
           reconciling_items,
           ["control", "from_account", "to_account", "amount", "lines",
            "kind", "recorded_by", "explanation"],
           [22, 40, 40, 16, 8, 24, 20, 90], {4})

    # ── Evidence ─────────────────────────────────────────────────────
    ws = _sheet(wb, "Evidence", "Documented cost by pool",
                "The proportion of each pool that has a document behind it.")
    _stamp(ws, period)
    _table(ws, 5, ["Pool", "Dollars", "Documented", "% documented"],
           evidence_coverage, ["pool", "dollars", "documented", "pct_documented"],
           [24, 18, 18, 16], {2, 3})

    # ── Exceptions ───────────────────────────────────────────────────
    ws = _sheet(wb, "Exceptions", "Where the standard bent",
                "Every departure, with the reason given at the time. An "
                "exception without a reason is the finding.")
    _stamp(ws, period)
    _table(ws, 5,
           ["Kind", "Subject", "Detail", "Amount", "Reason", "Who", "When"],
           exceptions,
           ["kind", "subject", "detail", "amount", "reason", "actor",
            "occurred_at"],
           [26, 34, 46, 16, 60, 20, 22], {4})

    # ── Rates ────────────────────────────────────────────────────────
    ws = _sheet(wb, "Rates", "Rates on file",
                "Each carries the seal of the decision set it came from. A "
                "database trigger refuses a rate whose seal does not match a "
                "sealed set.")
    _stamp(ws, period)
    last = _table(ws, 5,
                  ["Rate", "Pool", "Base", "Base amount", "Rate", "Status",
                   "Sealed by", "Seal"],
                  rates,
                  ["kind", "pool_amount", "base_type", "base_amount", "rate",
                   "status", "sealed_by", "seal_hash"],
                  [20, 16, 18, 16, 12, 14, 20, 66], {2, 4})
    _as_rate(ws, 6, last, 5)

    wb.save(out_path)
    return out_path


# ------------------------------------------------------------------- 990


def build_form_990(*, period: str, out_path: Path, functions: list[str],
                   categories: list[dict], totals: dict,
                   readiness: dict | None, documents: list[dict],
                   certification: dict | None = None) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    ws = _sheet(wb, "Part IX", "Statement of Functional Expenses",
                "Form 990 Part IX. Natural category down the side, function "
                "across the top, as the classification queue recorded it.")
    _stamp(ws, period)

    unallocated = Decimal(str((readiness or {}).get("unallocated") or 0))
    r = _caveat(ws, 5, certification_lines(certification))
    if unallocated:
        r = _caveat(ws, r, [
            "NOT FILEABLE — the allocation is incomplete.",
            f"{unallocated:,.2f} of expense has not been classified to a "
            f"function. It appears in its own column and is never spread "
            f"across the three the return prints: an allocation that "
            f"distributes unjudged cost is one nobody can support.",
            "The column totals are therefore short by that amount, on purpose.",
        ])

    # "Not applicable" is printed, not dropped. The handler has always
    # accumulated it and the sheet has never shown it, so a row whose total
    # exceeded its three functions gave the reader nothing to reconcile to —
    # and until `092` that was the entire payroll, $2,191,777.54 of it, on a
    # tax return. A column that is usually zero is cheaper than a total that
    # does not foot.
    headers = ["Natural category", "Lines", "Program", "Management and general",
               "Fundraising", "Not yet classified", "Not applicable", "Total"]
    keys = ["natural_category", "lines", "PROGRAM", "MANAGEMENT_AND_GENERAL",
            "FUNDRAISING", "NOT_YET_CLASSIFIED", "NOT_APPLICABLE", "total"]
    r = _table(ws, r, headers, categories, keys,
               [42, 10, 18, 24, 18, 22, 18, 18], {3, 4, 5, 6, 7, 8})

    r += 1
    ws.cell(row=r, column=1, value="Total").font = BOLD
    for i, key in enumerate(keys[2:], start=3):
        c = ws.cell(row=r, column=i, value=float(totals.get(key) or 0))
        c.font = BOLD
        c.number_format = "#,##0.00"

    # ── Attachments ──────────────────────────────────────────────────
    ws = _sheet(wb, "Attachments", "Documents on file",
                "What the return rests on. A document the register lists but "
                "cannot produce is a citation, not evidence.")
    _stamp(ws, period)
    _table(ws, 5,
           ["Document", "Kind", "Supports", "Received", "From", "Bytes",
            "SHA-256"],
           documents,
           ["filename", "kind", "supports", "received_at", "from_whom",
            "byte_size", "sha256"],
           [52, 26, 12, 22, 24, 14, 66])

    wb.save(out_path)
    return out_path
