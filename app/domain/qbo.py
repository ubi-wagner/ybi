"""
QuickBooks Online ingestion.

The thing that trips up QBO imports is that its exports are *reports*, not
tables. The General Ledger export is hierarchical: account section headers,
indented transaction rows, "Total for ..." subtotals, blank spacer rows, and
a running Balance column that is meaningless once the rows are extracted.
Column names also drift between QBO versions and between the "export to
Excel" and "export to CSV" paths.

So parsing is driven by a named ImportProfile rather than hardcoded column
positions. A new export shape is configuration, not a code change — and the
profile is stored with the import so a reviewer can see exactly how the file
was read two years later.

Every parse is a dry run until its control totals tie. Nothing reaches the
ledger tables on the strength of "it looked right".
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from .core import money

# QBO writes negatives three ways depending on report and locale.
_PARENS = re.compile(r"^\((.*)\)$")
_TOTAL_ROW = re.compile(r"^\s*total\s+(for\s+)?", re.I)
_BEGIN_BAL = re.compile(r"beginning\s+balance", re.I)

# QuickBooks closes a parent account with "Total for <parent> with sub-accounts".
# That row is a subtree rollup, not a leaf total, and reconciling it against
# lines coded directly to the parent will always fail — most parents carry no
# direct lines at all.
_ROLLUP_SUFFIX = re.compile(r"\s+with\s+sub-?accounts\s*$", re.I)


def qbo_amount(raw) -> Decimal:
    """QBO emits 1,234.56 / (1,234.56) / -1,234.56 / $1,234.56 / blank."""
    if raw is None:
        return Decimal(0)
    s = str(raw).strip().replace("$", "").replace(",", "").replace("\u00a0", "")
    if not s or s in {"-", "—"}:
        return Decimal(0)
    m = _PARENS.match(s)
    if m:
        s = "-" + m.group(1)
    try:
        return money(s)
    except Exception:
        return Decimal(0)


@dataclass
class ImportProfile:
    """A named, versioned recipe for reading one export shape.

    Column keys are matched case-insensitively against the detected header
    row, and several aliases are allowed per field because QBO renames things
    between versions ("Memo/Description" vs "Description" vs "Memo").
    """
    profile_id: str
    report: str                      # GENERAL_LEDGER | PROFIT_LOSS | BALANCE_SHEET | TIME_ACTIVITY
    aliases: dict[str, tuple[str, ...]] = field(default_factory=dict)
    debit_credit_columns: bool = False
    account_in_section_header: bool = True
    strip_totals: bool = True
    date_formats: tuple[str, ...] = ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y")

    def resolve(self, header: list[str]) -> dict[str, int]:
        norm = [(h or "").strip().lower() for h in header]
        out: dict[str, int] = {}
        for field_name, options in self.aliases.items():
            for opt in options:
                o = opt.strip().lower()
                if o in norm:
                    out[field_name] = norm.index(o)
                    break
        return out


QBO_GENERAL_LEDGER = ImportProfile(
    profile_id="qbo-gl-v1",
    report="GENERAL_LEDGER",
    aliases={
        "date": ("date", "txn date", "transaction date"),
        "txn_type": ("transaction type", "type"),
        "num": ("num", "no.", "number", "doc num"),
        "name": ("name", "vendor", "customer", "name/vendor/employee",
                 "payee", "customer/project"),
        "memo": ("memo/description", "description", "memo"),
        "split": ("split", "account"),
        "amount": ("amount",),
        "debit": ("debit",),
        "credit": ("credit",),
        "balance": ("balance", "running balance"),
        "class_": ("class",),
        "location": ("location",),
    },
)

QBO_PROFIT_LOSS = ImportProfile(
    profile_id="qbo-pl-v1",
    report="PROFIT_LOSS",
    aliases={"account": ("", "account", "distribution account"),
             "amount": ("total", "amount")},
    account_in_section_header=False,
)

QBO_TIME_ACTIVITY = ImportProfile(
    profile_id="qbo-time-v1",
    report="TIME_ACTIVITY",
    aliases={
        "date": ("date", "activity date"),
        "employee": ("employee", "name", "vendor/employee"),
        "customer": ("customer", "customer/project", "customer:job"),
        "service": ("service", "service item"),
        "hours": ("duration", "hours", "time"),
        "billable": ("billable", "billable?"),
        "memo": ("description", "memo", "notes"),
        "class_": ("class",),
    },
    account_in_section_header=False,
)


@dataclass
class StagedLine:
    row_number: int
    account: str
    date: str
    txn_type: str
    num: str
    name: str
    memo: str
    split: str
    amount: Decimal
    class_: str = ""
    location: str = ""
    customer_job: str = ""          # raw "Customer:Job" string if present
    objective_hint: str = ""        # the segment after the colon

    #: Which occurrence of an otherwise identical line this is. A journal can
    #: genuinely carry the same date, type, number, name, account, amount and
    #: memo twice — two $100 ticket sales, two identical splits of one entry —
    #: and they are different lines. Without this they hash alike, and the
    #: promote path's ON CONFLICT DO NOTHING silently keeps one of them: 58
    #: keys collided in the 2025 export, dropping 71 rows and $24,082.67 of
    #: activity with no error anywhere.
    occurrence: int = 0

    @property
    def natural_key(self) -> str:
        """Stable identity for a QBO line across re-exports. QBO does not
        expose transaction ids in report exports, so identity is composed
        from the fields that do not drift, plus which occurrence of that
        combination this is. Re-exporting the same period yields the same
        lines in the same order, so the occurrence is stable too."""
        import hashlib
        parts = "|".join([self.date, self.txn_type, self.num, self.name,
                          self.account, f"{self.amount:.2f}", self.memo[:80],
                          str(self.occurrence)])
        return hashlib.sha256(parts.encode()).hexdigest()[:24]


@dataclass
class StagedImport:
    profile_id: str
    source_name: str
    sha256: str
    lines: list[StagedLine] = field(default_factory=list)
    #: "Total for <account>" — leaf totals, keyed by qualified account path.
    subtotals: dict[str, Decimal] = field(default_factory=dict)
    #: "Total for <account> with sub-accounts" — subtree rollups, which
    #: reconcile against the account and everything beneath it.
    rollups: dict[str, Decimal] = field(default_factory=dict)
    #: "Beginning Balance" — the account's balance carried in from the prior
    #: year, keyed by qualified account path. Not a transaction, so it is not
    #: a line; but without it the general ledger cannot be tied to the balance
    #: sheet, because a balance sheet states a position and the ledger states
    #: a year of movement. Opening plus movement is the position.
    openings: dict[str, Decimal] = field(default_factory=dict)
    #: Accounts whose "Total for" row carried no printed figure — an export
    #: that saved formulas without their cached values. The subtotal control
    #: cannot be evaluated for these, and a control that cannot be evaluated
    #: must say so rather than compare against an assumed zero and pass.
    unprinted_subtotals: set[str] = field(default_factory=set)
    warnings: list[str] = field(default_factory=list)
    skipped: int = 0

    @property
    def total(self) -> Decimal:
        return money(sum(l.amount for l in self.lines))

    def account_totals(self) -> dict[str, Decimal]:
        out: dict[str, Decimal] = {}
        for l in self.lines:
            out[l.account] = money(out.get(l.account, Decimal(0)) + l.amount)
        return out

    def subtotal_check(self) -> list[tuple[str, Decimal, Decimal, Decimal]]:
        """QBO prints its own 'Total for <account>' rows. Reconciling parsed
        rows against them proves the parser did not drop or duplicate lines —
        an internal control that costs nothing and catches almost every
        parsing mistake."""
        parsed = self.account_totals()
        out = []
        for account, printed in self.subtotals.items():
            if account in self.unprinted_subtotals:
                continue        # nothing printed to reconcile against
            got = parsed.get(account, Decimal(0))
            out.append((account, printed, got, money(got - printed)))
        return sorted(out, key=lambda r: abs(r[3]), reverse=True)

    def subtree_total(self, account: str) -> Decimal:
        """Sum of an account and every sub-account beneath it."""
        prefix = account + ":"
        return money(sum(
            (l.amount for l in self.lines
             if l.account == account or l.account.startswith(prefix)),
            Decimal(0)))

    def rollup_check(self) -> list[tuple[str, Decimal, Decimal, Decimal]]:
        """Reconcile the "with sub-accounts" rollups against their subtrees.

        This is the control that proves the account hierarchy was reconstructed
        correctly, not merely that individual leaves add up."""
        out = []
        for account, printed in self.rollups.items():
            got = self.subtree_total(account)
            out.append((account, printed, got, money(got - printed)))
        return sorted(out, key=lambda r: abs(r[3]), reverse=True)


def _read_rows(path: Path) -> list[list[str]]:
    if path.suffix.lower() in {".xlsx", ".xls", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        return [["" if c is None else str(c) for c in row]
                for row in ws.iter_rows(values_only=True)]
    with io.open(path, encoding="utf-8-sig", newline="") as f:
        return [row for row in csv.reader(f)]


def _find_header(rows: list[list[str]], profile: ImportProfile) -> tuple[int, dict[str, int]]:
    """QBO puts a company name, report title and date range above the header,
    and sometimes a blank spacer column at position 0. Scan for the first row
    that resolves enough known columns."""
    need = {"date", "amount"} if not profile.debit_credit_columns else {"date", "debit"}
    for i, row in enumerate(rows[:40]):
        cols = profile.resolve(row)
        if need.issubset(cols.keys()):
            return i, cols
        if "date" in cols and ("amount" in cols or "debit" in cols):
            return i, cols
    raise ValueError(
        "Could not locate a header row. QBO exports vary; add the column "
        "names from this file to an ImportProfile rather than editing code."
    )


def parse_general_ledger(path: Path, profile: ImportProfile = QBO_GENERAL_LEDGER,
                         sha256: str = "") -> StagedImport:
    rows = _read_rows(path)
    hdr_idx, cols = _find_header(rows, profile)
    staged = StagedImport(profile_id=profile.profile_id, source_name=path.name, sha256=sha256)

    def cell(row: list[str], key: str) -> str:
        i = cols.get(key)
        return (row[i].strip() if i is not None and i < len(row) and row[i] else "")

    # QuickBooks prints the account tree as header / total pairs in column 0
    # rather than by indentation, so the qualified path has to be tracked as a
    # stack. Without it, sub-accounts sharing a leaf name under different
    # parents collide — YBI's 2025 chart has "Drive AM" under both
    # "3900 Grant Income" and "Grant Expenses", and merging them silently
    # overstates the account by the whole of the other one.
    #
    # One further wrinkle: QuickBooks prints a parent's own direct lines and
    # its "Total for <parent>" *before* the sub-account sections, then closes
    # with "Total for <parent> with sub-accounts". So a header following a
    # leaf total may be either that account's child or its sibling, and column
    # 0 does not say which. The rollup rows do: an account is a parent exactly
    # when it has a "with sub-accounts" total. Collect those first, then use
    # them to decide whether a leaf total closes its frame.
    parents = {
        _ROLLUP_SUFFIX.sub("", _TOTAL_ROW.sub("", (r[0] or "").strip())).strip()
        for r in rows[hdr_idx + 1:]
        if r and (r[0] or "").strip()
        and _TOTAL_ROW.match((r[0] or "").strip())
        and _ROLLUP_SUFFIX.search((r[0] or "").strip())
    }

    stack: list[str] = []
    for n, row in enumerate(rows[hdr_idx + 1:], start=hdr_idx + 2):
        if not any((c or "").strip() for c in row):
            continue

        first = (row[0] or "").strip() if row else ""

        # "Total for 5010 Depreciation Expense" — capture as a control, skip row
        if first and _TOTAL_ROW.match(first):
            label = _TOTAL_ROW.sub("", first).strip()
            is_rollup = bool(_ROLLUP_SUFFIX.search(label))
            account = _ROLLUP_SUFFIX.sub("", label).strip()
            raw_amt = cell(row, "amount")
            amt = qbo_amount(raw_amt)
            # An export that saved formulas without their cached values leaves
            # this cell empty. Falling back to the last non-empty cell on the
            # row picks up the running Balance column — opening plus movement,
            # not movement — and the subtotal then reconciles against the
            # wrong figure. Record the absence instead; a control that cannot
            # be evaluated is not a control that passed.
            printed = bool(raw_amt.strip())
            if account:
                if account in stack:
                    depth = len(stack) - 1 - stack[::-1].index(account)
                    path = ":".join(stack[:depth + 1])
                    # A leaf total closes the frame only when the account has
                    # no sub-accounts; otherwise its children are still to
                    # come and the frame stays open until the rollup.
                    if is_rollup or account not in parents:
                        del stack[depth:]
                else:
                    path = ":".join(stack + [account])
                if is_rollup:
                    staged.rollups[path] = amt
                else:
                    staged.subtotals[path] = amt
                if not printed:
                    staged.unprinted_subtotals.add(path)
            continue

        # "Beginning Balance" — not a transaction, but not noise either: it is
        # the position the account was carried in at. QuickBooks prints it
        # under the account header with the figure in the running Balance
        # column and nothing in Amount, so it carries no date and would
        # otherwise be swallowed by the dateless-row guard below.
        #
        # It is what makes the ledger tieable to the balance sheet at all: a
        # balance sheet states a position, a ledger states a year of movement,
        # and only opening plus movement is a position.
        # Different exports put the label in different columns — under the
        # distribution account in one, under the date in another — so the
        # test is the label plus the absence of an amount. A transaction
        # always carries an amount; an opening position never does.
        if (any(_BEGIN_BAL.search(c or "") for c in row)
                and not cell(row, "amount")
                and not cell(row, "debit") and not cell(row, "credit")):
            if stack:
                staged.openings[":".join(stack)] = qbo_amount(cell(row, "balance"))
            continue

        # Section header: something in column 0 and no date on the row
        if profile.account_in_section_header and first and not cell(row, "date"):
            stack.append(first)
            continue

        date = cell(row, "date")
        if not date:
            staged.skipped += 1
            continue

        if profile.debit_credit_columns:
            amount = money(qbo_amount(cell(row, "debit")) - qbo_amount(cell(row, "credit")))
        else:
            amount = qbo_amount(cell(row, "amount"))

        account = ":".join(stack) if stack else cell(row, "split")
        if not account:
            staged.warnings.append(f"row {n}: no account context, line held for review")

        name = cell(row, "name")
        customer_job = name if ":" in name else ""
        hint = name.split(":")[-1].strip() if ":" in name else ""

        staged.lines.append(StagedLine(
            row_number=n, account=account, date=date,
            txn_type=cell(row, "txn_type"), num=cell(row, "num"),
            name=name, memo=cell(row, "memo"), split=cell(row, "split"),
            amount=amount, class_=cell(row, "class_"),
            location=cell(row, "location"),
            customer_job=customer_job, objective_hint=hint,
        ))

    # Number the repeats so identical lines stay distinct. In file order,
    # which is the order the same export produces every time.
    seen: dict[str, int] = {}
    repeated = 0
    for l in staged.lines:
        base = l.natural_key            # occurrence is still 0 here
        n = seen.get(base, 0)
        if n:
            l.occurrence = n
            repeated += 1
        seen[base] = n + 1
    if repeated:
        staged.warnings.append(
            f"{repeated} line(s) repeat an otherwise identical line — same "
            f"date, type, num, name, account, amount and memo. Genuine split "
            f"lines and genuine repeats both look like this, so each is kept "
            f"as its own row rather than collapsed.")

    collisions = len(staged.lines) - len({l.natural_key for l in staged.lines})
    if collisions:
        staged.warnings.append(
            f"{collisions} line(s) still share a natural key after numbering "
            f"the repeats. These would be lost on promote; do not accept.")
    return staged


def parse_time_activity(path: Path, profile: ImportProfile = QBO_TIME_ACTIVITY) -> list[dict]:
    """QBO Time Activity detail. This is the feed that makes labour
    distribution contemporaneous rather than reconstructed — the single
    highest-value import in the system."""
    rows = _read_rows(path)
    hdr_idx = 0
    cols: dict[str, int] = {}
    for i, row in enumerate(rows[:40]):
        c = profile.resolve(row)
        if "employee" in c and "hours" in c:
            hdr_idx, cols = i, c
            break
    if not cols:
        raise ValueError("Time activity header not found; extend QBO_TIME_ACTIVITY aliases.")

    out = []
    for row in rows[hdr_idx + 1:]:
        def g(k):
            i = cols.get(k)
            return (row[i].strip() if i is not None and i < len(row) and row[i] else "")
        if not g("employee") or not g("date"):
            continue
        cust = g("customer")
        out.append({
            "date": g("date"),
            "employee": g("employee"),
            "customer_job": cust,
            "objective_hint": cust.split(":")[-1].strip() if ":" in cust else cust,
            "service": g("service"),
            "hours": qbo_amount(g("hours")),
            "billable": g("billable").lower() in {"yes", "true", "y", "1"},
            "memo": g("memo"),
            "class_": g("class_"),
        })
    return out


def objective_hints(staged: StagedImport) -> dict[str, Decimal]:
    """Customer:Job already encodes the cost objective in YBI's file — 188
    revenue rows worth $2.03M carry it. Surface it as a classification
    proposal rather than making anyone retype it."""
    out: dict[str, Decimal] = {}
    for l in staged.lines:
        if l.objective_hint:
            out[l.objective_hint] = money(out.get(l.objective_hint, Decimal(0)) + l.amount)
    return dict(sorted(out.items(), key=lambda kv: abs(kv[1]), reverse=True))


# ---------------------------------------------------------------------
# Profit and loss
# ---------------------------------------------------------------------

#: The four sections a QBO P&L prints, in the order it prints them. The
#: section is what decides whether an account is a cost, a revenue or neither
#: — never the account number. YBI's 2025 chart has revenue sitting in the
#: 5xxx expense range (5107 Interest Income, 5108 Other Income) and income in
#: the net-assets range (3991 MBAC), so number-based inference is wrong here.
PL_SECTIONS = ("Income", "COGS", "Expense", "Other Income")

_PL_SECTION_HEADERS = {
    "income": "Income",
    "cost of goods sold": "COGS",
    "expenses": "Expense",
    "other income": "Other Income",
}


@dataclass
class ProfitLoss:
    """A parsed QuickBooks Profit and Loss.

    ``accounts`` maps the qualified account path to its section and amount.
    ``leaf_section`` maps the bare leaf name to its section, which is what the
    general ledger needs in order to tell a cost account from a balance sheet
    account.
    """

    source_name: str
    sha256: str = ""
    accounts: dict[str, tuple[str, Decimal]] = field(default_factory=dict)
    rollups: dict[str, tuple[str, Decimal]] = field(default_factory=dict)
    section_totals: dict[str, Decimal] = field(default_factory=dict)
    net_income: Decimal = Decimal(0)
    warnings: list[str] = field(default_factory=list)

    @property
    def income_total(self) -> Decimal:
        return self.section_totals.get("Income", Decimal(0))

    @property
    def cogs_total(self) -> Decimal:
        return self.section_totals.get("COGS", Decimal(0))

    @property
    def expense_total(self) -> Decimal:
        return self.section_totals.get("Expense", Decimal(0))

    @property
    def other_income_total(self) -> Decimal:
        return self.section_totals.get("Other Income", Decimal(0))

    @property
    def leaf_section(self) -> dict[str, str]:
        out: dict[str, str] = {}
        for path, (section, _) in self.accounts.items():
            out[path.split(":")[-1]] = section
        for path, (section, _) in self.rollups.items():
            out.setdefault(path.split(":")[-1], section)
        return out

    def check_net_income(self) -> Decimal:
        """Variance between the printed net income and the sections."""
        derived = (self.income_total - self.cogs_total - self.expense_total
                   + self.other_income_total)
        return money(derived - self.net_income)


def parse_profit_loss(path: Path, sha256: str = "") -> ProfitLoss:
    """Parse a QBO Profit and Loss export.

    Same header/total hierarchy as the general ledger, but with section
    headers (Income, Cost of Goods Sold, Expenses, Other Income) framing the
    account tree and a single amount column.
    """
    rows = _read_rows(path)
    pl = ProfitLoss(source_name=path.name, sha256=sha256)

    # On a P&L every "Total for X" is a subtree total, so every one of them
    # names a parent. Requiring the "with sub-accounts" suffix here — as the
    # general ledger does — left this set empty, nothing was ever pushed, and
    # every account landed at the top level. Accounts sharing a leaf name
    # across sections then collided: Drive AM, Digital Engineering, DLA Grant
    # and Youth Entrepreneurship each exist in both Income and Expenses, and
    # the later one silently overwrote the earlier.
    parents = {
        _ROLLUP_SUFFIX.sub("", _TOTAL_ROW.sub("", (r[0] or "").strip())).strip()
        for r in rows
        if r and (r[0] or "").strip()
        and _TOTAL_ROW.match((r[0] or "").strip())
    }

    def amount_of(row: list[str]) -> Decimal:
        for c in reversed(row[1:]):
            if (c or "").strip():
                return qbo_amount(c)
        return Decimal(0)

    section = ""
    stack: list[str] = []
    for row in rows:
        if not row or not any((c or "").strip() for c in row):
            continue
        label = (row[0] or "").strip()
        if not label:
            continue

        low = label.lower()
        if low in _PL_SECTION_HEADERS:
            section = _PL_SECTION_HEADERS[low]
            stack = []
            continue

        if _TOTAL_ROW.match(label):
            name = _ROLLUP_SUFFIX.sub(
                "", _TOTAL_ROW.sub("", label).strip()).strip()
            amt = amount_of(row)

            low_name = name.lower()
            if low_name in _PL_SECTION_HEADERS:
                pl.section_totals[_PL_SECTION_HEADERS[low_name]] = amt
                stack = []
                continue

            # Every "Total for X" on a P&L is a subtree rollup, whether or not
            # it carries the "with sub-accounts" suffix. Recording it as an
            # account would double count: 4029 Sponsorships totals 270,209 and
            # its three children total the same 270,209.
            if name in stack:
                depth = len(stack) - 1 - stack[::-1].index(name)
                qualified = ":".join(stack[:depth + 1])
                del stack[depth:]
            else:
                qualified = ":".join(stack + [name])
            pl.rollups[qualified] = (section, amt)
            continue

        if low.startswith("net income"):
            pl.net_income = amount_of(row)
            continue
        # Computed lines, not accounts. "Net Other Income" repeats the Other
        # Income section total and would double it.
        if low.startswith(("net operating income", "net other income",
                           "gross profit")):
            continue

        # An account line. A parent may carry its own amount AND have children
        # beneath it — 4000 Contributions Income books 85,820.23 directly and
        # still has sub-accounts — so the parent's own amount is recorded and
        # the parent is then pushed, rather than one or the other.
        if not section:
            continue
        amt = amount_of(row)
        has_amount = any((c or "").strip() for c in row[1:])
        qualified = ":".join(stack + [label])
        if has_amount:
            pl.accounts[qualified] = (section, amt)
        if label in parents:
            stack.append(label)

    return pl


# ---------------------------------------------------------------------------
# Balance sheet
# ---------------------------------------------------------------------------

#: The rows that switch which side of the sheet we are on. Everything else
#: that carries no amount is a group header, and QuickBooks gives every group
#: a matching "Total for" row.
_BS_SIDES = {
    "assets": "ASSET",
    "liabilities": "LIABILITY",
    "equity": "EQUITY",
}


@dataclass
class BalanceSheet:
    """A parsed QuickBooks Balance Sheet.

    Two controls matter and both are printed on the face of the report, which
    is what makes them worth checking rather than trusting: the sheet has to
    balance, and the net income it carries has to be the net income on the
    Profit and Loss. A balance sheet that disagrees with the P&L means one of
    the two exports is from a different moment, and everything built on either
    is suspect.
    """

    source_name: str
    sha256: str = ""
    #: qualified path -> (side, amount)
    accounts: dict[str, tuple[str, Decimal]] = field(default_factory=dict)
    rollups: dict[str, tuple[str, Decimal]] = field(default_factory=dict)
    net_income: Decimal = Decimal(0)
    warnings: list[str] = field(default_factory=list)

    def _rollup(self, *names: str) -> Decimal:
        """A named total, by its leaf name, wherever it sits in the tree."""
        for path, (_, amount) in self.rollups.items():
            if path.split(":")[-1].lower() in {n.lower() for n in names}:
                return amount
        return Decimal(0)

    @property
    def assets(self) -> Decimal:
        return self._rollup("Assets")

    @property
    def liabilities(self) -> Decimal:
        return self._rollup("Liabilities")

    @property
    def equity(self) -> Decimal:
        return self._rollup("Equity")

    @property
    def fixed_assets(self) -> Decimal:
        return self._rollup("Fixed Assets")

    def subtotal_checks(self) -> list[tuple[str, Decimal, Decimal, Decimal]]:
        """Every printed total against the sum of what sits directly under it.

        The general ledger parser learned this the expensive way: a tree
        rebuilt from header/total pairs can look entirely reasonable while an
        account hangs off the wrong parent. An arithmetic check on every
        subtotal is what makes the shape provable rather than plausible.

        Returns (path, printed, derived, variance) per rollup.
        """
        out = []
        for path, (_, printed) in self.rollups.items():
            prefix = path + ":"
            depth = path.count(":") + 1
            # A parent may carry its own balance as well as children — the
            # credit-card accounts do, with the card itself holding a balance
            # and each cardholder a sub-account. That own balance sits at the
            # parent's own path and belongs inside the parent's total.
            derived = self.accounts.get(path, ("", Decimal(0)))[1]
            for p, (_, amt) in self.accounts.items():
                # An account that has a total of its own is represented by
                # that total, not by its own balance; counting both is how the
                # same $298.51 of card balances appeared twice.
                if (p.startswith(prefix) and p.count(":") == depth
                        and p not in self.rollups):
                    derived += amt
            for p, (_, amt) in self.rollups.items():
                if p.startswith(prefix) and p.count(":") == depth:
                    derived += amt
            out.append((path, money(printed), money(derived),
                        money(derived - printed)))
        return out

    def failing_subtotals(self) -> list[tuple[str, Decimal, Decimal, Decimal]]:
        return [c for c in self.subtotal_checks() if c[3] != 0]

    def check_balance(self) -> Decimal:
        """Assets less liabilities and equity. Zero, or the export is wrong."""
        return money(self.assets - self.liabilities - self.equity)

    def check_net_income(self, pl_net_income: Decimal) -> Decimal:
        """The cross-statement tie. The two reports have to be the same
        moment in the same books."""
        return money(self.net_income - pl_net_income)


def parse_balance_sheet(path: Path, sha256: str = "") -> BalanceSheet:
    """Parse a QBO Balance Sheet export.

    The same header/total shape as the other reports, with three wrinkles this
    sheet actually contains:

    * A parent may carry its own amount inline and still have children —
      ``2100 Payroll Liabilities  $0.00`` followed by six tax accounts.
    * A parent may have no total row at all. ``1530 Computer Equipment``
      exists only to hold ``1560 Equipment (Parent)``; treating it as an
      account would invent a line, and failing to push it would flatten its
      children into the wrong parent.
    * Contra accounts are negative and belong with what they offset:
      accumulated depreciation sits beside the cost it reduces, which is what
      makes the fixed-asset section answer the 200.436(b) question at all.
    """
    rows = _read_rows(path)
    bs = BalanceSheet(source_name=path.name, sha256=sha256)

    # Every name that gets a "Total for" row. A header without one is a label
    # QuickBooks printed and never closed — "1530 Computer Equipment" is there
    # to introduce "1560 Equipment (Parent)" and is never totalled. Pushing it
    # as a parent means nothing ever pops it, and the next sibling falls
    # inside: 1570 TBB5, an $8.9M building, filed under computer equipment.
    # The subtotal control below is what catches this if it ever returns.
    totalled = {
        _ROLLUP_SUFFIX.sub("", _TOTAL_ROW.sub("", (r[0] or "").strip())).strip()
        for r in rows
        if r and (r[0] or "").strip()
        and _TOTAL_ROW.match((r[0] or "").strip())
    }

    def amount_of(row: list[str]) -> Decimal | None:
        for c in reversed(row[1:]):
            if (c or "").strip():
                return qbo_amount(c)
        return None

    side = ""
    started = False
    stack: list[str] = []
    for row in rows:
        if not row or not any((c or "").strip() for c in row):
            continue
        label = (row[0] or "").strip()
        if not label:
            continue
        low = label.lower()

        # Report furniture. The sheet proper starts at "Assets"; everything
        # above it is the company name, the report name and the as-of date.
        if not started:
            if low == "assets":
                started = True
            else:
                continue
        if low.startswith(("accrual basis", "cash basis", "balance sheet",
                           "as of")) or low == "total":
            continue

        if _TOTAL_ROW.match(label):
            name = _ROLLUP_SUFFIX.sub(
                "", _TOTAL_ROW.sub("", label).strip()).strip()
            amt = amount_of(row) or Decimal(0)
            if name in stack:
                depth = len(stack) - 1 - stack[::-1].index(name)
                qualified = ":".join(stack[:depth + 1])
                del stack[depth:]
            else:
                qualified = ":".join(stack + [name])
            bs.rollups[qualified] = (side, amt)
            continue

        if low.startswith("net income"):
            bs.net_income = amount_of(row) or Decimal(0)
            # Also an equity account in its own right: leaving it out of the
            # tree makes Equity fail to foot by exactly the year's result,
            # which looks like a parser bug and is one.
            bs.accounts[":".join(stack + [label])] = (side, bs.net_income)
            continue

        amt = amount_of(row)

        # A group header. No amount of its own, so it can only be a parent.
        if amt is None:
            if low == "liabilities and equity":
                # The root the other two sit inside. Without it, "Total for
                # Liabilities and Equity" has nothing beneath it to derive
                # from and the sheet's own footing check cannot be proved.
                stack = [label]
                continue
            if low in _BS_SIDES:
                side = _BS_SIDES[low]
                stack = ([stack[0], label] if stack and
                         stack[0].lower() == "liabilities and equity"
                         else [label])
                continue
            if label in totalled:
                stack.append(label)
            else:
                bs.warnings.append(
                    f"group header {label!r} has no total row; treated as a "
                    f"label rather than a parent")
            continue

        # A leaf, or a parent carrying its own balance. Either way the amount
        # is this account's own; a parent's subtree arrives as its "Total for".
        qualified = ":".join(stack + [label])
        if qualified in bs.accounts:
            bs.warnings.append(f"duplicate account path {qualified}")
        bs.accounts[qualified] = (side, amt)

        # QuickBooks prints a parent's own balance before its children, so a
        # parent with an inline amount still has to go on the stack. It is a
        # parent exactly when a "Total for" bearing its name follows.
        if label in totalled:
            stack.append(label)

    return bs
