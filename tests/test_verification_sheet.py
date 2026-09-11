"""The worksheet and the document must describe the same eighteen things.

`docs/FOR_TOM_TO_VERIFY.md` is the narrative — how each discrepancy was found
and why it matters. `scripts/verification_sheet.py` turns the same items into a
workbook and a printable worksheet.

Two copies of one list is exactly the shape this codebase keeps getting bitten
by: coverage computed in two places answered 13.0% and 2.2% at the same moment,
and a hand-kept map of what each screen calls was wrong four times in one run.
The narrative genuinely has to be prose and the worksheet genuinely has to be
structured, so the two cannot be derived from each other — which leaves testing
that they agree.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from scripts.verification_sheet import ITEMS, STATUS, build_workbook

DOC = Path(__file__).resolve().parent.parent / "docs" / "FOR_TOM_TO_VERIFY.md"


def refs_in_document() -> set[str]:
    """The items, not the groups they sit under.

    `## 1. Depreciation` is a heading over six items; `### 1.1` is one of
    them. Only the second kind is a thing to settle — and the worked example
    at `## 0.` is an item that happens to have no group above it.
    """
    text = DOC.read_text()
    return (set(re.findall(r"^###\s+(\d+\.\d+)\s", text, re.M))
            | set(re.findall(r"^##\s+(0)\.\s", text, re.M)))


def test_the_worksheet_and_the_document_carry_the_same_items():
    sheet = {i.ref for i in ITEMS}
    doc = refs_in_document()
    assert sheet == doc, (
        f"only in the worksheet: {sorted(sheet - doc)}; "
        f"only in the document: {sorted(doc - sheet)}"
    )


def test_the_count_the_document_claims_is_the_count_it_has():
    """The first line says "Eighteen things". A document that miscounts itself
    is one a reader stops trusting on everything else."""
    words = {"Sixteen": 16, "Seventeen": 17, "Eighteen": 18,
             "Nineteen": 19, "Twenty": 20, "Twenty-one": 21,
             "Twenty-two": 22, "Twenty-three": 23, "Twenty-four": 24}
    head = DOC.read_text().split("---")[0]
    # Longest first: "Twenty" is a substring of "Twenty-one", and matching it
    # would read a document that says twenty-one as claiming twenty — a test
    # miscounting the document it is checking for miscounting itself.
    claimed = next((n for w, n in sorted(words.items(), key=lambda kv: -len(kv[0]))
                    if w in head), None)
    assert claimed is not None, "the opening does not say how many there are"
    assert claimed == len(ITEMS)


def test_every_item_asks_something_and_says_what_it_moves():
    for item in ITEMS:
        assert item.asks.strip(), f"{item.ref} asks nothing"
        assert item.moves.strip(), f"{item.ref} does not say what turns on it"
        assert item.found.strip(), f"{item.ref} does not say what was found"
        assert item.figure.strip(), f"{item.ref} carries no figure"


def test_exactly_one_item_is_already_answered():
    """The Bacon credit is the worked example. A second pre-filled row would
    read as a partly-completed sheet rather than as a pattern to follow."""
    answered = [i for i in ITEMS if i.answered]
    assert len(answered) == 1 and answered[0].ref == "0"
    assert answered[0].answered[0] in STATUS


def test_the_status_choices_are_not_a_yes_or_no():
    """"I have checked and the record is right" and "I have corrected it" lead
    to different work at our end, and "somebody else has to answer this" is a
    real outcome that otherwise looks like silence."""
    assert len(STATUS) >= 4
    joined = " ".join(STATUS).upper()
    for needed in ("CONFIRMED", "CORRECTED", "STILL CHECKING", "SOMEBODY ELSE"):
        assert needed in joined


def test_the_workbook_builds_and_offers_the_statuses():
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(build_workbook()))
    assert {"Start here", "Verify", "Choices"} <= set(wb.sheetnames)
    verify = wb["Verify"]
    assert verify.max_row == len(ITEMS) + 1
    offered = {c.value for row in wb["Choices"].iter_rows() for c in row
               if c.value}
    assert set(STATUS) <= offered
    # The columns we filled in are locked; the ones they answer are not.
    headings = [c.value for c in verify[1]]
    ref_col = headings.index("Ref") + 1
    answer_col = headings.index("Your answer") + 1
    assert verify.cell(row=2, column=ref_col).protection.locked
    assert not verify.cell(row=2, column=answer_col).protection.locked
