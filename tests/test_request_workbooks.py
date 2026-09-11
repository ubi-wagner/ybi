"""The ask, and the imperfect answer that comes back.

Every workbook that returns will have something wrong with it. These tests
are mostly about that: what happens to the eleven bad cells in a four hundred
row register, and whether what happens to them is something a person can act
on in a minute.
"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.domain.request_forms import (ASSET_REGISTER, FORMS, PEOPLE_ROSTER,
                                      SPACE_INVENTORY, Kind)
from app.domain.request_intake import (WorkbookNotRecognised,
                                       read_request_workbook)
from app.domain.request_workbook import (IDENTITY_CELL, IDENTITY_SHEET,
                                         build_request_workbook)


def fill(form, rows, *, period="2025", known_rows=None, mangle=None):
    """Issue a workbook, type `rows` into it, and read it back.

    Deliberately round-trips through real bytes rather than calling the
    parser on a dict: the thing being tested is a file that went out, was
    opened in Excel, and came back.
    """
    data = build_request_workbook(form, period, known_rows=known_rows or [],
                                  blank_rows=len(rows) + 5)
    wb = load_workbook(BytesIO(data))
    ws = wb[form.sheet]
    at = {}
    for i, cell in enumerate(ws[1], start=1):
        text = str(cell.value or "").strip().rstrip(" *")
        if text:
            at[text] = i
    start = 3 + len(known_rows or [])
    for r, row in enumerate(rows, start=start):
        for heading, value in row.items():
            ws.cell(row=r, column=at[heading], value=value)
    if mangle:
        mangle(wb)
    out = BytesIO()
    wb.save(out)
    return read_request_workbook(out.getvalue(), form)


# ── The form is defined once ──────────────────────────────────────────

@pytest.mark.parametrize("form", list(FORMS.values()), ids=list(FORMS))
def test_every_column_the_workbook_asks_for_can_be_read_back(form):
    """The writer and the reader run off one definition, so there is no way
    for a workbook to ask a question the parser cannot accept — which is the
    shape that produced two coverage figures six-fold apart."""
    filled = fill(form, [])
    assert filled.missing_columns == ()
    assert filled.form is form


@pytest.mark.parametrize("form", list(FORMS.values()), ids=list(FORMS))
def test_the_workbook_says_what_it_is(form):
    data = build_request_workbook(form, "2025")
    wb = load_workbook(BytesIO(data))
    line = wb[IDENTITY_SHEET][IDENTITY_CELL].value
    assert form.name in line and f"v{form.version}" in line and "2025" in line


@pytest.mark.parametrize("form", list(FORMS.values()), ids=list(FORMS))
def test_every_choice_column_offers_its_choices(form):
    """A free-text funding source is how "EDA grant (I think)" ends up in the
    column the allowability test reads."""
    data = build_request_workbook(form, "2025")
    wb = load_workbook(BytesIO(data))
    for col in form.columns:
        if col.kind is not Kind.CHOICE:
            continue
        assert "Choices" in wb.sheetnames
        offered = {c.value for row in wb["Choices"].iter_rows() for c in row}
        assert set(col.choices) <= offered


# ── A blank is not a zero ─────────────────────────────────────────────

def test_a_blank_is_unanswered_and_a_zero_is_an_answer():
    """The intake version of "unclassified cost is never defaulted into a
    pool". Somebody who has checked and found no federal money has told us
    something; somebody who skipped the column has not, and chasing the
    second is the whole point of asking."""
    filled = fill(ASSET_REGISTER, [
        {"Asset ID or tag": "A-1", "Description": "Checked, none federal",
         "Original cost, before any reimbursement": 1000,
         "Federal money in it": 0},
        {"Asset ID or tag": "A-2", "Description": "Nobody looked",
         "Original cost, before any reimbursement": 2000},
    ])
    checked, skipped = filled.rows
    assert checked.values["federal_amount"] == Decimal("0")
    assert skipped.values["federal_amount"] is None
    assert filled.answered("federal_amount") == 1


def test_a_blank_yes_no_is_not_a_no():
    filled = fill(ASSET_REGISTER, [
        {"Asset ID or tag": "A-1", "Description": "x",
         "Original cost, before any reimbursement": 1,
         "Was this counted as cost share?": "no"},
        {"Asset ID or tag": "A-2", "Description": "y",
         "Original cost, before any reimbursement": 1},
    ])
    assert filled.rows[0].values["counted_as_cost_share"] is False
    assert filled.rows[1].values["counted_as_cost_share"] is None


# ── The imperfect answer ──────────────────────────────────────────────

def test_one_bad_cell_costs_one_cell_not_the_file():
    filled = fill(ASSET_REGISTER, [
        {"Asset ID or tag": "A-1", "Description": "Fine",
         "Original cost, before any reimbursement": 1000,
         "Placed in service": "2019-03-01"},
        {"Asset ID or tag": "A-2", "Description": "Vague date",
         "Original cost, before any reimbursement": 2000,
         "Placed in service": "2019ish"},
        {"Asset ID or tag": "A-3", "Description": "Also fine",
         "Original cost, before any reimbursement": 3000},
    ])
    assert len(filled.rows) == 3
    assert len(filled.usable) == 3, "a bad date must not lose the asset"
    assert filled.total("gross_cost") == Decimal("6000")
    problem = next(p for p in filled.problems if p.heading == "Placed in service")
    assert problem.row == 4 and problem.value == "2019ish"
    assert "is not a date" in problem.says
    assert 'row 4' in str(problem)


def test_money_as_a_person_types_it():
    filled = fill(ASSET_REGISTER, [
        {"Asset ID or tag": "A-1", "Description": "Commas and a symbol",
         "Original cost, before any reimbursement": "$1,234.56"},
        {"Asset ID or tag": "A-2", "Description": "Accounting negative",
         "Original cost, before any reimbursement": "(500.00)"},
    ])
    assert filled.problems == []
    assert filled.rows[0].values["gross_cost"] == Decimal("1234.56")
    assert filled.rows[1].values["gross_cost"] == Decimal("-500.00")


def test_a_sentence_where_a_number_belongs_is_refused_not_mined():
    """"about 40,000" is a sentence. Reaching in for the digits would put a
    figure nobody typed into a rate."""
    filled = fill(ASSET_REGISTER, [
        {"Asset ID or tag": "A-1", "Description": "x",
         "Original cost, before any reimbursement": "about 40,000"},
    ])
    assert filled.rows[0].values["gross_cost"] is None
    assert any("is not a number" in p.says for p in filled.problems)


def test_a_choice_outside_the_list_names_the_list():
    filled = fill(SPACE_INVENTORY, [
        {"Building": "Tech Block 5", "Suite or area": "210",
         "Usable square feet": 1200, "What it is used for": "rented out",
         "Occupied or empty": "OCCUPIED", "Who is in it": "Acme"},
    ])
    problem = next(p for p in filled.problems if p.heading == "What it is used for")
    assert "TENANT" in problem.says and "PROGRAM" in problem.says


def test_a_choice_is_read_however_it_is_capitalised():
    filled = fill(SPACE_INVENTORY, [
        {"Building": "Tech Block 5", "Suite or area": "210",
         "Usable square feet": 1200, "What it is used for": "tenant",
         "Occupied or empty": "occupied", "Who is in it": "Acme"},
    ])
    assert filled.problems == []
    assert filled.rows[0].values["use"] == "TENANT"


def test_a_row_missing_what_it_cannot_be_recorded_without_is_held_back():
    filled = fill(ASSET_REGISTER, [
        {"Description": "No identifier, no cost",
         "Serial or VIN": "SN-9"},
    ])
    assert filled.usable == []
    assert len(filled.incomplete) == 1
    problem = next(p for p in filled.problems if p.row == 3)
    assert "held back rather than half written" in problem.says


def test_a_prefilled_row_nobody_touched_is_a_queue_not_a_fault():
    """The six balance-sheet rows come back untouched if nobody got to them.
    Reporting those as incomplete would bury the rows somebody actually
    started and left, which are the ones worth a phone call."""
    filled = fill(ASSET_REGISTER, [
        {"Asset ID or tag": "A-1", "Description": "Done properly",
         "Original cost, before any reimbursement": 10},
    ], known_rows=[{"gl_account": "1501 Tech Block Bldg 5"},
                   {"gl_account": "1511 Buildings"}])
    assert len(filled.untouched) == 2
    assert len(filled.incomplete) == 0
    assert len(filled.usable) == 1


# ── Things people do to spreadsheets ──────────────────────────────────

def test_a_column_moved_does_not_shift_the_data_along_one():
    """Reading by position turns an inserted column into silently wrong data
    in every column to the right of it."""
    def insert(wb):
        wb[ASSET_REGISTER.sheet].insert_cols(2)
        wb[ASSET_REGISTER.sheet].cell(row=1, column=2, value="Our own notes")

    filled = fill(ASSET_REGISTER, [
        {"Asset ID or tag": "A-1", "Description": "Still lines up",
         "Original cost, before any reimbursement": 4242},
    ], mangle=insert)
    assert filled.problems == []
    assert filled.rows[0].values["description"] == "Still lines up"
    assert filled.rows[0].values["gross_cost"] == Decimal("4242")


def test_a_deleted_column_is_named_rather_than_read_as_blank():
    def drop(wb):
        ws = wb[ASSET_REGISTER.sheet]
        for i, cell in enumerate(ws[1], start=1):
            if str(cell.value or "").startswith("Federal money"):
                ws.delete_cols(i)
                break

    filled = fill(ASSET_REGISTER, [
        {"Asset ID or tag": "A-1", "Description": "x",
         "Original cost, before any reimbursement": 1},
    ], mangle=drop)
    assert "Federal money in it" in filled.missing_columns
    assert any("older copy of the form" in p.says for p in filled.problems)


def test_deleting_the_explanation_row_does_not_lose_the_first_answer():
    """Somebody tidies row 2 away and starts typing there. Assuming the data
    begins on row 3 would drop their first asset without a word."""
    def tidy(wb):
        wb[ASSET_REGISTER.sheet].delete_rows(2)

    filled = fill(ASSET_REGISTER, [
        {"Asset ID or tag": "A-1", "Description": "First one",
         "Original cost, before any reimbursement": 7},
        {"Asset ID or tag": "A-2", "Description": "Second",
         "Original cost, before any reimbursement": 8},
    ], mangle=tidy)
    assert [r.values["asset_id"] for r in filled.rows] == ["A-1", "A-2"]


def test_blank_rows_between_answers_are_not_answers():
    filled = fill(ASSET_REGISTER, [
        {"Asset ID or tag": "A-1", "Description": "x",
         "Original cost, before any reimbursement": 1},
        {},
        {"Asset ID or tag": "A-2", "Description": "y",
         "Original cost, before any reimbursement": 2},
    ])
    assert len(filled.rows) == 2


# ── The wrong file ────────────────────────────────────────────────────

def test_the_wrong_form_says_which_form_it_is():
    data = build_request_workbook(SPACE_INVENTORY, "2025")
    with pytest.raises(WorkbookNotRecognised) as exc:
        read_request_workbook(data, ASSET_REGISTER)
    assert "SPACE_INVENTORY" in str(exc.value)
    assert "ASSET_REGISTER" in str(exc.value)


def test_a_grid_pasted_into_a_new_file_is_refused_with_a_reason():
    """The likeliest wrong file of all: somebody copies the sheet into a
    fresh workbook and sends that."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = ASSET_REGISTER.sheet
    wb.active["A1"] = "Asset ID or tag"
    out = BytesIO()
    wb.save(out)
    with pytest.raises(WorkbookNotRecognised) as exc:
        read_request_workbook(out.getvalue(), ASSET_REGISTER)
    assert IDENTITY_SHEET in str(exc.value)


def test_the_roster_reads_a_name_and_an_address():
    filled = fill(PEOPLE_ROSTER, [
        {"Payroll ID": "E-014", "First name": "Dolores",
         "Working email address": "dwallace@ybi.org", "Still employed?": "yes"},
    ], known_rows=[])
    assert filled.problems == []
    row = filled.rows[0]
    assert row.values["email"] == "dwallace@ybi.org"
    assert row.values["still_employed"] is True
    assert row.usable
